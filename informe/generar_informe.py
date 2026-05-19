"""
generar_informe.py — Construye el informe técnico del compilador ENIGMA
en formato .docx (Word).

Incluye todas las secciones solicitadas:
  1. Portada
  2. Introducción
  3. Descripción del compilador
  4. Integración del compilador
  5. Pruebas realizadas (léxicas, sintácticas, semánticas, manejo de errores)
  6. Problemas encontrados
  7. Conclusiones

Cada prueba incluye: código probado, resultado esperado, resultado obtenido
y captura de pantalla real generada por el script `generar_capturas.py`.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor


ROOT_DIR = Path(__file__).resolve().parent.parent
INFORME_DIR = ROOT_DIR / "informe"
SCREENSHOTS_DIR = INFORME_DIR / "screenshots"
LOGO_PATH = ROOT_DIR / "Logo ENIGMA.png"

OUTPUT_PATH = INFORME_DIR / "Informe_Tecnico_Compilador_ENIGMA.docx"


# ───────────────────────── helpers de estilo ───────────────────────────

def set_cell_shading(cell, color_hex: str) -> None:
    """Aplica color de fondo a una celda."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def add_horizontal_line(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "0F4C81")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)
        run.font.name = "Calibri"


def add_paragraph(doc: Document, text: str, *, bold: bool = False,
                  italic: bool = False, size: int = 11,
                  align=WD_ALIGN_PARAGRAPH.JUSTIFY) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.runs[0] if p.runs else p.add_run("")
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(11)


def add_code_block(doc: Document, code: str) -> None:
    """Inserta un bloque de código monoespaciado en una tabla 1×1."""
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F4F6F8")
    # Quitar margen extra de la tabla
    cell.paragraphs[0].text = ""
    for linea in code.splitlines():
        p = cell.add_paragraph()
        run = p.add_run(linea if linea else "\u00a0")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    # Eliminar el primer párrafo vacío
    first = cell.paragraphs[0]
    first._element.getparent().remove(first._element)
    doc.add_paragraph()


def add_screenshot(doc: Document, nombre: str, descripcion: str,
                   ancho_cm: float = 16.5) -> None:
    ruta = SCREENSHOTS_DIR / nombre
    if not ruta.exists():
        add_paragraph(doc, f"[Captura faltante: {nombre}]", italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(ruta), width=Cm(ancho_cm))

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = caption.add_run(f"Figura — {descripcion}")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_paragraph()


def add_prueba(doc: Document, *, titulo: str, codigo: str, esperado: str,
               obtenido: str, captura: str, captura_desc: str) -> None:
    add_heading(doc, titulo, level=3)

    add_paragraph(doc, "Código probado:", bold=True)
    add_code_block(doc, codigo)

    add_paragraph(doc, "Resultado esperado:", bold=True)
    add_paragraph(doc, esperado)

    add_paragraph(doc, "Resultado obtenido:", bold=True)
    add_paragraph(doc, obtenido)

    add_paragraph(doc, "Evidencia (captura del compilador):", bold=True)
    add_screenshot(doc, captura, captura_desc)


# ─────────────────────────── PORTADA ───────────────────────────────────

def construir_portada(doc: Document) -> None:
    # Espacio superior
    for _ in range(2):
        doc.add_paragraph()

    if LOGO_PATH.exists():
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.add_run().add_picture(str(LOGO_PATH), width=Cm(6.5))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("UNIVERSIDAD MARIANO GÁLVEZ DE GUATEMALA")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Facultad de Ingeniería en Sistemas de Información")
    r.font.size = Pt(13)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Curso: Compiladores")
    r.bold = True
    r.font.size = Pt(14)

    doc.add_paragraph()
    add_horizontal_line(doc.add_paragraph())

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("INFORME TÉCNICO")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Compilador para el Lenguaje de Seguridad Empresarial "ENIGMA"')
    r.bold = True
    r.italic = True
    r.font.size = Pt(15)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Tema: Integración, ejecución y pruebas de las fases léxica, sintáctica y semántica."
    )
    r.font.size = Pt(11)
    r.italic = True

    add_horizontal_line(doc.add_paragraph())
    doc.add_paragraph()
    doc.add_paragraph()

    # Integrantes en tabla
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Integrantes del grupo")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)

    tabla = doc.add_table(rows=4, cols=2)
    tabla.alignment = WD_ALIGN_PARAGRAPH.CENTER
    integrantes = [
        ("Jordi García",   "Analizador Léxico"),
        ("Germán Archila", "Analizador Sintáctico y Semántico — Controlador"),
        ("Víctor Morales", "Interfaz Gráfica (PyQt6)"),
        ("Juan Carlos",    "Integración, pruebas y documentación"),
    ]
    for fila, (nombre, modulo) in enumerate(integrantes):
        c1 = tabla.cell(fila, 0)
        c2 = tabla.cell(fila, 1)
        c1.text = nombre
        c2.text = modulo
        for c in (c1, c2):
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for parag in c.paragraphs:
                parag.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in parag.runs:
                    run.font.size = Pt(11)
                    run.font.name = "Calibri"

    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fecha = date.today().strftime("%d de %B de %Y")
    meses = {
        "January": "enero", "February": "febrero", "March": "marzo",
        "April": "abril",  "May": "mayo",        "June": "junio",
        "July": "julio",   "August": "agosto",   "September": "septiembre",
        "October": "octubre", "November": "noviembre", "December": "diciembre",
    }
    for en, es in meses.items():
        fecha = fecha.replace(en, es)
    r = p.add_run(f"Guatemala, {fecha}")
    r.font.size = Pt(12)
    r.italic = True

    doc.add_page_break()


# ─────────────────── 2. INTRODUCCIÓN ───────────────────────────────────

def construir_introduccion(doc: Document) -> None:
    add_heading(doc, "1. Introducción", level=1)

    add_paragraph(doc,
        "El presente informe describe el desarrollo, integración y validación del "
        "compilador ENIGMA, un sistema construido en Python 3 que procesa el "
        "“Lenguaje de Control de Accesos Empresarial”. Este lenguaje, diseñado por "
        "el grupo, permite expresar de manera declarativa políticas de seguridad "
        "basadas en roles (RBAC) y atributos (ABAC) que en un entorno real podrían "
        "regir el acceso a los módulos de una empresa (ventas, inventario, reportes, "
        "facturación, etc.).")

    add_paragraph(doc,
        "El compilador toma como entrada un archivo fuente con extensión .acl, "
        "lo somete a las tres fases clásicas de análisis —léxico, sintáctico y "
        "semántico— y produce como salida una tabla de tokens, un árbol de "
        "sintaxis abstracta (AST), una tabla de símbolos con las entidades "
        "declaradas y, opcionalmente, un archivo JSON con las políticas compiladas "
        "(Fase 4). Cuando el código contiene errores, el compilador los detecta, "
        "los reporta con línea y columna, y continúa el análisis para identificar "
        "varios errores en una sola ejecución gracias al “Modo Pánico”.")

    add_heading(doc, "Objetivo del proyecto", level=2)
    add_paragraph(doc,
        "Construir un compilador completo y funcional, con interfaz gráfica "
        "didáctica, que permita comprender en profundidad cómo se integran las "
        "fases de un compilador. El objetivo académico es doble: (a) aplicar los "
        "conceptos teóricos del curso a un caso real (control de accesos empresarial) "
        "y (b) demostrar la robustez del sistema mediante pruebas exhaustivas que "
        "ejerciten tanto los casos válidos como todos los tipos de error posibles.")

    add_heading(doc, "Importancia de las pruebas y la integración", level=2)
    add_paragraph(doc,
        "Un compilador es un software compuesto por módulos altamente interdependientes: "
        "si el lexer entrega un flujo de tokens inconsistente, el parser falla; si el "
        "parser produce un AST mal formado, el analizador semántico no podrá validar "
        "las reglas. Por ello, la integración no es opcional sino crítica. Las pruebas, "
        "por su parte, son la única forma objetiva de garantizar que cada fase respeta "
        "la gramática del lenguaje y reacciona correctamente ante entradas inválidas. "
        "En este proyecto se ejecutaron pruebas léxicas, sintácticas y semánticas, "
        "tanto con casos correctos como con catálogos completos de errores, validando "
        "que el compilador se mantiene estable y produce diagnósticos útiles.")

    doc.add_page_break()


# ─────────────── 3. DESCRIPCIÓN DEL COMPILADOR ─────────────────────────

def construir_descripcion(doc: Document) -> None:
    add_heading(doc, "2. Descripción del compilador", level=1)

    add_heading(doc, "Funcionalidades principales", level=2)
    funciones = [
        "Análisis léxico con reconocimiento de palabras reservadas (case-insensitive), "
        "identificadores, números enteros y decimales, cadenas, operadores y delimitadores.",
        "Detección de 6 tipos de errores léxicos: carácter foráneo, cadena sin cerrar, "
        "número mal formado, palabra reservada mal escrita (con sugerencias por distancia "
        "de Levenshtein), identificador que inicia con dígito e identificador con guion bajo.",
        "Análisis sintáctico descendente recursivo (recursive-descent) basado en la "
        "gramática BNF del lenguaje, con precedencia para los operadores lógicos (O < Y < No).",
        "Recuperación de errores sintácticos en Modo Pánico: el parser se sincroniza sobre "
        "los delimitadores “;” y “}” para seguir analizando y reportar varios errores en "
        "una sola pasada.",
        "Análisis semántico con 7 reglas (ERR_SEM_01 a ERR_SEM_07): unicidad, integridad "
        "referencial, compatibilidad de tipos, dominio RBAC, expresiones booleanas, "
        "operaciones entre tipos y conflicto de políticas.",
        "Tabla de Símbolos visual con entidades (Rol, Usuario, Módulo), variables ABAC "
        "globales pre-inyectadas (Horario, MontoVenta, UbicacionIP) y políticas asociadas.",
        "Interfaz gráfica PyQt6 con editor con resaltado de sintaxis, Error Lens, vista "
        "de árbol del AST y vista gráfica circular animada.",
        "Modo Didáctico: animación paso a paso de las tres fases sincronizada con el "
        "código fuente.",
        "Live Compile: compilación silenciosa con debounce mientras se escribe el código.",
        "Exportación a JSON de las políticas compiladas (Fase 4).",
    ]
    for f in funciones:
        add_bullet(doc, f)

    add_heading(doc, "Lenguaje utilizado", level=2)
    add_paragraph(doc,
        "El compilador está escrito íntegramente en Python 3.12. La interfaz gráfica "
        "se implementa con PyQt6 y las pruebas automatizadas se ejecutan con pytest. "
        "El lenguaje fuente analizado —ENIGMA— es un DSL (Domain-Specific Language) "
        "case-insensitive con palabras en español como “Definir”, “Rol”, “Usuario”, "
        "“Permitir”, “Denegar”, “Si”, “Entonces”, “Mientras”, “Elegir”, “Intentar”, "
        "“Atrapar”, entre otras.")

    add_heading(doc, "Módulos implementados", level=2)
    modulos = [
        ("lexer/",     "tokens.py, lexer.py, error_handler.py — Motor léxico."),
        ("parser/",    "ast_nodes.py, parser.py — Construcción del AST."),
        ("semantic/",  "semantic_analyzer.py, symbol_table.py, semantic_errors.py — Validaciones."),
        ("gui/",       "main_window.py, code_editor.py, token_table.py, ast_graph_widget.py, "
                       "symbol_table_widget.py, icons.py — Interfaz gráfica."),
        ("codegen/",   "policy_exporter.py — Exportación de políticas a JSON (Fase 4)."),
        ("controller.py", "Coordinador del pipeline y del Modo Didáctico."),
        ("examples/",  "Archivos .acl con casos válidos y catálogos de errores."),
        ("test/",      "Suite automatizada de pruebas con pytest."),
    ]
    for nombre, desc in modulos:
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(nombre + " — ")
        r1.bold = True
        r1.font.name = "Consolas"
        r1.font.size = Pt(11)
        r2 = p.add_run(desc)
        r2.font.name = "Calibri"
        r2.font.size = Pt(11)

    add_heading(doc, "Proceso de compilación", level=2)
    add_paragraph(doc,
        "El flujo se ejecuta de forma estrictamente secuencial. Si una fase produce "
        "errores que comprometen a la siguiente, el pipeline se detiene de manera "
        "controlada y la GUI muestra exactamente en qué fase se abortó.")
    fases = [
        "Fase 1 — Léxico: el módulo Lexer recorre el código fuente carácter por "
        "carácter mediante expresiones regulares priorizadas y emite una lista de "
        "objetos Token. Cualquier carácter inesperado se reporta al ErrorHandler "
        "sin detener la tokenización.",
        "Fase 2 — Sintáctico: el Parser consume la lista de tokens y construye un "
        "AST formado por 14 tipos de nodos. Si encuentra un error, activa el Modo "
        "Pánico y continúa.",
        "Fase 3 — Semántico: el SemanticAnalyzer recorre el AST con patrón Visitor, "
        "alimentando una Tabla de Símbolos. Aplica las 7 validaciones del lenguaje.",
        "Fase 4 — Generación (opcional): si las tres fases anteriores son exitosas, "
        "el PolicyExporter convierte el AST y la Tabla de Símbolos en un documento "
        "JSON listo para ser consumido por un sistema RBAC/ABAC externo.",
    ]
    for f in fases:
        add_bullet(doc, f)

    add_paragraph(doc,
        "La siguiente captura muestra la interfaz principal del compilador con un "
        "programa válido cargado y todas las fases ejecutadas con éxito.")
    add_screenshot(doc, "01_programa_completo_tokens.png",
                   "Vista general del compilador ENIGMA con la tabla de tokens activa.")

    doc.add_page_break()


# ─────────────── 4. INTEGRACIÓN DEL COMPILADOR ─────────────────────────

def construir_integracion(doc: Document) -> None:
    add_heading(doc, "3. Integración del compilador", level=1)

    add_heading(doc, "Cómo se unieron los módulos", level=2)
    add_paragraph(doc,
        "La integración se realizó alrededor de un único punto de orquestación: "
        "la clase Controller. Esta clase actúa como mediador entre la interfaz "
        "gráfica y los analizadores. Cuando el usuario pulsa “Analizar” (Ctrl+Enter) "
        "o escribe en el editor (Live Compile), el Controller ejecuta el método "
        "_run_pipeline(), que:")
    pasos = [
        "Toma el texto del editor (gui.main_window.MainWindow.get_code()).",
        "Lo entrega al Lexer y obtiene la lista de tokens y los errores léxicos.",
        "Si hay errores léxicos, los pinta con Error Lens en el editor y aborta.",
        "En caso contrario, pasa los tokens al Parser que produce el AST y los errores sintácticos.",
        "Si hay errores sintácticos, los muestra en el panel y aborta la fase semántica.",
        "Si todo está limpio, el AST se pasa al SemanticAnalyzer, que llena la Tabla "
        "de Símbolos y reporta cualquier violación.",
        "Finalmente, la GUI refresca tres pestañas independientes: Tokens, AST Gráfico "
        "y Tabla de Símbolos, además del Panel de Errores cuando aplica.",
    ]
    for p in pasos:
        add_bullet(doc, p)

    add_paragraph(doc,
        "Cada equipo desarrolló su módulo por separado siguiendo contratos claros "
        "(tipos de retorno y estructuras de datos comunes definidas en lexer/tokens.py "
        "y parser/ast_nodes.py). La integración se llevó a cabo en tres sesiones de "
        "trabajo: integración léxico-sintáctica, integración con la GUI y, por último, "
        "incorporación del analizador semántico y de la tabla de símbolos visual.")

    add_heading(doc, "Problemas encontrados durante la integración", level=2)
    problemas = [
        "Inconsistencias en el case-sensitivity: el lexer normaliza palabras "
        "reservadas a minúsculas, pero el parser inicialmente las comparaba "
        "respetando mayúsculas. Se solucionó normalizando lexemas en el lexer.",
        "Coordenadas (línea, columna) desactualizadas tras inserciones de tokens "
        "virtuales en la recuperación del parser. Se ajustó propagando el "
        "start_index y end_index reales del token previo.",
        "Refrescos visuales redundantes cuando el Live Compile disparaba el "
        "pipeline en cada pulsación. Se introdujo un QTimer con debounce de 350 ms.",
        "El visor gráfico del AST se quedaba dibujando una versión antigua tras "
        "cambiar de archivo. Se añadió clear_graph() antes de cada análisis.",
        "La Tabla de Símbolos no mostraba las variables ABAC globales hasta que "
        "se invocaba el analizador semántico. Se decidió inyectarlas en el "
        "constructor de SymbolTable para que estuvieran disponibles desde el inicio.",
    ]
    for p in problemas:
        add_bullet(doc, p)

    add_heading(doc, "Soluciones aplicadas", level=2)
    add_paragraph(doc,
        "Para resolver de forma definitiva los choques de integración se adoptaron "
        "tres prácticas: (1) introducción de pruebas automatizadas con pytest que "
        "verifican el contrato de cada fase de manera aislada; (2) revisión del "
        "Controller como único punto de entrada al pipeline, evitando que la GUI "
        "instancie analizadores directamente; (3) uso de tipos de datos inmutables "
        "(dataclasses) para los tokens y los nodos del AST, lo que evita estados "
        "compartidos accidentales entre fases.")

    add_paragraph(doc,
        "La siguiente captura corresponde al sistema integrado mostrando el AST "
        "construido a partir del archivo programa_completo.acl. Allí se puede "
        "verificar que las tres fases trabajan en cadena: tokens generados, AST "
        "renderizado y análisis semántico sin errores.")
    add_screenshot(doc, "02_programa_completo_ast.png",
                   "AST gráfico del programa completo (sistema integrado).")
    add_screenshot(doc, "03_programa_completo_simbolos.png",
                   "Tabla de Símbolos generada al final del pipeline.")

    doc.add_page_break()


# ───────────────────── 5. PRUEBAS REALIZADAS ────────────────────────────

def construir_pruebas(doc: Document) -> None:
    add_heading(doc, "4. Pruebas realizadas", level=1)

    add_paragraph(doc,
        "Las pruebas se diseñaron en dos grandes bloques: (a) casos válidos que el "
        "compilador debe aceptar sin reportar ningún error y (b) catálogos de errores "
        "que ejercitan todos los diagnósticos implementados. Cada prueba se ejecutó "
        "directamente desde la GUI cargando los archivos .acl ubicados en la carpeta "
        "examples/ del proyecto, y se capturó la pantalla resultante para servir "
        "como evidencia.")

    # ─────────── 4.1 Pruebas léxicas ────────────────────────────────
    add_heading(doc, "4.1 Pruebas léxicas", level=2)

    add_prueba(doc,
        titulo="Prueba L-01 — Tokenización de un programa válido",
        codigo=(
            "Definir Rol Gerente;\n"
            "Definir Usuario Ana;\n"
            "Definir Modulo Ventas;\n"
            "Usuario Ana = Rol Gerente;\n"
            "Rol Gerente = Permitir Consultar Ventas;"
        ),
        esperado=(
            "El lexer debe emitir una lista de tokens correctamente clasificados "
            "(palabras reservadas, identificadores y delimitadores) sin reportar "
            "ningún error léxico."
        ),
        obtenido=(
            "El compilador generó la tabla de tokens completa para el archivo "
            "empresa_ventas.acl: cada palabra fue clasificada según TipoToken con "
            "su línea y columna. No se reportaron errores léxicos."
        ),
        captura="04_empresa_ventas_tokens.png",
        captura_desc="Tokens del archivo empresa_ventas.acl (caso válido).",
    )

    add_prueba(doc,
        titulo="Prueba L-02 — Catálogo completo de errores léxicos",
        codigo=(
            "Definir Rol @Gerente;          // ERR_LEX_01\n"
            'Mostrar "Cadena nunca cerrada  // ERR_LEX_02\n'
            "Mostrar 15.5.2;                 // ERR_LEX_03\n"
            'Defini Rol Asistente;           // ERR_LEX_04\n'
            "Definir Rol 123Gerente;         // ERR_LEX_05\n"
            "Definir Rol mi_rol;             // ERR_LEX_06"
        ),
        esperado=(
            "El lexer debe detectar los 6 tipos de error léxico, marcarlos con "
            "subrayado ondulado en el editor (Error Lens) y poblar el Panel de "
            "Errores con línea, columna y mensaje. Para “Defini” debe sugerir "
            "“Definir” por distancia de Levenshtein."
        ),
        obtenido=(
            "Se detectaron todos los errores esperados (uno por cada categoría) "
            "y el panel inferior derecho mostró el listado con los códigos "
            "ERR_LEX_01 a ERR_LEX_06. El compilador no colapsa: continúa "
            "escaneando hasta el final del archivo."
        ),
        captura="07_errores_lexicos.png",
        captura_desc="Detección masiva de errores léxicos.",
    )

    # ─────────── 4.2 Pruebas sintácticas ──────────────────────────────
    add_heading(doc, "4.2 Pruebas sintácticas", level=2)

    add_prueba(doc,
        titulo="Prueba S-01 — Construcción correcta del AST",
        codigo=(
            "Si Verdadero Entonces {\n"
            "    Permitir Consultar Reportes;\n"
            "}\n"
            "Mientras Verdadero {\n"
            "    Validar Ventas;\n"
            "}"
        ),
        esperado=(
            "El parser debe construir un AST con un ProgramNode raíz, un "
            "SiEntoncesNode y un MientrasNode con sus respectivos bloques."
        ),
        obtenido=(
            "El visor gráfico del AST renderizó la estructura jerárquica con "
            "nodos circulares para programa_completo.acl: condiciones, bucles, "
            "bloque Elegir/Caso/Terminar e Intentar/Atrapar quedaron representados "
            "correctamente."
        ),
        captura="02_programa_completo_ast.png",
        captura_desc="AST gráfico construido sin errores sintácticos.",
    )

    add_prueba(doc,
        titulo="Prueba S-02 — Modo Pánico con múltiples errores sintácticos",
        codigo=(
            "Definir Rol Asistente        // falta ';'\n"
            'Definir Usuario "Empleado101"; // literal en vez de identificador\n'
            "Si Verdadero Entonces {        // falta '}'\n"
            "    Permitir Consultar Inventario;\n"
            "Asignar Rol Auditor;           // keyword inválida\n"
            "Denegar Permitir Reportes;     // doble acción"
        ),
        esperado=(
            "El parser debe detectar todos los errores, sincronizarse sobre los "
            "delimitadores “;” y “}” y reportar cada problema con su línea y "
            "una explicación clara, sin colapsar."
        ),
        obtenido=(
            "El compilador reportó múltiples errores sintácticos en una sola "
            "ejecución, demostrando la efectividad del Modo Pánico. La fase "
            "semántica fue bloqueada correctamente, tal como indica la barra de "
            "estado de la GUI."
        ),
        captura="08_errores_sintacticos.png",
        captura_desc="Modo Pánico recuperándose de múltiples errores sintácticos.",
    )

    # ─────────── 4.3 Pruebas semánticas ──────────────────────────────
    add_heading(doc, "4.3 Pruebas semánticas", level=2)

    add_prueba(doc,
        titulo="Prueba SE-01 — Programa empresarial completo (caso válido)",
        codigo=(
            "Definir Rol Gerente;\n"
            "Definir Rol Cajero;\n"
            "Definir Usuario Ana;\n"
            "Definir Modulo Ventas;\n"
            "Usuario Ana = Rol Gerente;\n"
            "Rol Gerente = Permitir Consultar Ventas;\n"
            "Rol Cajero  = Permitir Registrar Ventas;"
        ),
        esperado=(
            "Las 7 reglas semánticas deben pasar sin errores. La Tabla de "
            "Símbolos debe registrar las entidades con su tipo y políticas."
        ),
        obtenido=(
            "El analizador semántico reportó cero errores. La Tabla de Símbolos "
            "mostró Gerente, Cajero, Ana y Ventas con sus tipos correctos, "
            "junto con las variables ABAC globales (Horario, MontoVenta, "
            "UbicacionIP) inyectadas automáticamente."
        ),
        captura="06_empresa_ventas_simbolos.png",
        captura_desc="Tabla de Símbolos del caso empresarial válido.",
    )

    add_prueba(doc,
        titulo="Prueba SE-02 — Redeclaración de entidad (ERR_SEM_01)",
        codigo=(
            "Definir Rol Gerente;\n"
            "Definir Usuario Gerente;  // El identificador ya existe como Rol"
        ),
        esperado=(
            "El analizador semántico debe detectar la redeclaración del "
            "identificador “Gerente” y abortar la fase semántica con un mensaje "
            "claro (ERR_SEM_01)."
        ),
        obtenido=(
            "El compilador reportó ERR_SEM_01 con la línea exacta del conflicto, "
            "explicando que el identificador ya estaba registrado como Rol. La "
            "Tabla de Símbolos refleja únicamente las entidades válidas declaradas "
            "antes del error."
        ),
        captura="09_errores_semanticos.png",
        captura_desc="Validación semántica: redeclaración detectada.",
    )

    # ─────────── 4.4 Manejo de errores ────────────────────────────────
    add_heading(doc, "4.4 Manejo de errores y casos límite", level=2)

    add_paragraph(doc,
        "Adicional a las pruebas anteriores, se ejecutó la batería completa "
        "incluida en examples/errores_semanticos.acl para validar los 7 errores "
        "semánticos del catálogo (ERR_SEM_01 a ERR_SEM_07). La política “Fail-Fast” "
        "del analizador detiene el análisis en el primer error encontrado, lo que "
        "obliga a comentar bloques para probar cada uno por separado. En todas las "
        "ejecuciones el compilador respondió de manera estable, sin lanzar "
        "excepciones no controladas.")

    add_paragraph(doc,
        "Pruebas adicionales realizadas (no mostradas aquí por brevedad pero "
        "documentadas en test/test_case.py):")
    extras = [
        "Programa vacío: la barra de estado indica que no hay nada para analizar.",
        "Programa con sólo comentarios: el lexer ignora correctamente las líneas "
        "que inician con “//”.",
        "Códigos extremadamente largos (>500 líneas) procesados por Live Compile "
        "sin congelar la interfaz, gracias al debounce.",
        "Cambio rápido de archivo: la GUI limpia las tres pestañas antes del "
        "siguiente análisis.",
    ]
    for e in extras:
        add_bullet(doc, e)

    add_screenshot(doc, "05_empresa_ventas_ast.png",
                   "AST del caso empresarial real (tienda Tech-Store).")

    doc.add_page_break()


# ─────────────────── 6. PROBLEMAS ENCONTRADOS ───────────────────────────

def construir_problemas(doc: Document) -> None:
    add_heading(doc, "5. Problemas encontrados", level=1)

    add_heading(doc, "Errores detectados durante el desarrollo", level=2)
    errores = [
        ("Tokens con coordenadas (línea, columna) incorrectas tras un salto de "
         "línea dentro de una cadena no cerrada.",
         "Se ajustó el contador de línea del Lexer para incrementarlo aun cuando "
         "el escaneo esté en estado de error."),
        ("El Parser entraba en bucle al encontrar un token inesperado al inicio "
         "de una sentencia.",
         "Se reforzó el método _sincronizar() del Modo Pánico para consumir al "
         "menos un token antes de continuar."),
        ("El analizador semántico fallaba con una excepción cuando el AST contenía "
         "nodos huérfanos por un parser que abortó.",
         "Se añadió una guarda al inicio de SemanticAnalyzer.analizar() que "
         "verifica que el árbol esté bien formado y se decidió bloquear la fase "
         "semántica si hay errores sintácticos."),
        ("La vista gráfica del AST se quedaba con nodos antiguos si se analizaba "
         "un archivo nuevo.",
         "Se invoca ast_graph.clear_graph() antes de cada análisis y al pulsar "
         "el botón Limpiar."),
        ("Las variables ABAC globales (Horario, MontoVenta, UbicacionIP) aparecían "
         "duplicadas en la Tabla de Símbolos al re-ejecutar el pipeline.",
         "Se reescribió SymbolTable.reset() para regenerar el contenido inicial "
         "en lugar de añadir variables."),
    ]
    for problema, solucion in errores:
        p = doc.add_paragraph(style="List Number")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run("Problema: ")
        run.bold = True
        p.add_run(problema)
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p2.paragraph_format.left_indent = Cm(0.75)
        r = p2.add_run("Solución: ")
        r.bold = True
        p2.add_run(solucion)

    add_heading(doc, "Dificultades durante el desarrollo", level=2)
    dificultades = [
        "Coordinar el trabajo paralelo de cuatro integrantes con responsabilidades "
        "distintas (lexer, parser/semántico, GUI e integración) requirió definir "
        "interfaces estables desde el inicio.",
        "PyQt6 obliga a ejecutar todo el código gráfico en el hilo principal, lo "
        "que dificultó probar el pipeline con scripts. Se resolvió creando un "
        "modo de captura automatizada que abre la ventana, ejecuta el análisis y "
        "guarda screenshots en disco.",
        "La gramática del lenguaje creció durante el proyecto (bloques Elegir, "
        "Intentar/Atrapar, condiciones lógicas con precedencia). Cada nueva "
        "estructura obligó a ampliar el AST, el parser y el visor gráfico.",
        "El diseño visual de la Tabla de Símbolos cambió varias veces hasta lograr "
        "mostrar adecuadamente entidades, sub-tipos, rol vinculado y políticas en "
        "un mismo widget legible.",
    ]
    for d in dificultades:
        add_bullet(doc, d)

    doc.add_page_break()


# ─────────────────────── 7. CONCLUSIONES ────────────────────────────────

def construir_conclusiones(doc: Document) -> None:
    add_heading(doc, "6. Conclusiones", level=1)

    conclusiones = [
        "El compilador ENIGMA cumple satisfactoriamente con las tres fases clásicas "
        "del análisis (léxico, sintáctico y semántico) y agrega una cuarta etapa de "
        "exportación a JSON, demostrando que es posible llevar un proyecto académico "
        "hasta un producto utilizable como motor de políticas RBAC/ABAC.",

        "La integración modular alrededor de una clase Controller demostró ser la "
        "decisión arquitectónica más acertada: permitió que cada integrante desarrollara "
        "su capa de forma independiente y que la unión final se realizara mediante "
        "contratos claros (Token, AST y SymbolTable), reduciendo significativamente los "
        "conflictos en la integración.",

        "Las pruebas con catálogos completos de errores (léxicos, sintácticos y "
        "semánticos) confirmaron la robustez del compilador. En todos los escenarios "
        "incorrectos el sistema reportó diagnósticos claros, conservó la estabilidad "
        "y no colapsó, gracias a la recuperación con Modo Pánico y a la política "
        "Fail-Fast del analizador semántico.",

        "La interfaz gráfica en PyQt6 con Live Compile, Error Lens y Modo Didáctico "
        "transformó al compilador en una herramienta pedagógica además de funcional, "
        "permitiendo observar paso a paso cómo se construyen los tokens, el AST y la "
        "Tabla de Símbolos.",

        "El proyecto evidenció la importancia de las pruebas automatizadas: la suite "
        "con pytest sirvió como red de seguridad cada vez que se modificó la gramática "
        "o se añadió una nueva regla semántica, evitando regresiones difíciles de "
        "detectar manualmente.",

        "El aprendizaje más valioso fue comprobar que un compilador no es un único "
        "programa monolítico, sino una cadena de transformaciones donde cada fase "
        "depende de la calidad de la anterior. Diseñar interfaces estables entre "
        "módulos y mantener una disciplina de pruebas continuas son la diferencia "
        "entre un compilador frágil y uno realmente confiable.",
    ]
    for i, c in enumerate(conclusiones, 1):
        p = doc.add_paragraph(style="List Number")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(c)
        run.font.name = "Calibri"
        run.font.size = Pt(11)

    doc.add_paragraph()
    add_horizontal_line(doc.add_paragraph())
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("— Fin del informe técnico —")
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


# ────────────────────────────── main ────────────────────────────────────

def main() -> None:
    doc = Document()

    # Márgenes y estilo base
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    estilo_normal = doc.styles["Normal"]
    estilo_normal.font.name = "Calibri"
    estilo_normal.font.size = Pt(11)

    construir_portada(doc)
    construir_introduccion(doc)
    construir_descripcion(doc)
    construir_integracion(doc)
    construir_pruebas(doc)
    construir_problemas(doc)
    construir_conclusiones(doc)

    doc.save(str(OUTPUT_PATH))
    print(f"[OK] Informe guardado en: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
